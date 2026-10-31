"""Build Nhan's shareable coordination memo using the bundled document runtime."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

root = Path(__file__).resolve().parents[2]
output = root.parent / 'deliverables' / 'KLEER_Nhan_Phoi_hop_va_Phe_duyet.docx'
output.parent.mkdir(parents=True, exist_ok=True)
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin = section.bottom_margin = Cm(1.8)
section.left_margin = section.right_margin = Cm(2)
for style_name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'List Bullet']:
    style = doc.styles[style_name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0,0,0)
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(7)
    style.paragraph_format.line_spacing = 1.12
doc.styles['Title'].font.size = Pt(21)
doc.styles['Heading 1'].font.size = Pt(17)
doc.styles['Heading 2'].font.size = Pt(12)
doc.styles['Heading 2'].font.bold = True
for name in ['Title','Subtitle','Heading 1','Heading 2']:
    doc.styles[name].font.color.rgb = RGBColor(0,0,0)
    doc.styles[name].paragraph_format.keep_with_next = True
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.core_properties.title = 'KLEER Phối hợp hoàn thành các task của Nhân'
doc.core_properties.author = 'Bùi Hiếu Nhân'
doc.core_properties.subject = 'Đề nghị phê duyệt và phối hợp theo từng thành viên'

def p(text, style=None):
    return doc.add_paragraph(text, style)

def bullets(items):
    for text in items: p(text, 'List Bullet')

def heading(text): doc.add_heading(text, level=2)

doc.add_paragraph('KLEER Phối hợp hoàn thành\ncác task của Nhân', 'Title')
p('Người gửi Bùi Hiếu Nhân\nNgày cập nhật 26 tháng 10 năm 2026')
p('Mình đã chuẩn bị và kiểm thử phần backend Quiz, taxonomy và tài liệu API. Để đóng các task còn lại, mình cần nhóm duyệt nghiệp vụ G3, cung cấp môi trường có WooCommerce và dữ liệu thật, rồi phối hợp kiểm thử tích hợp. Các phần sau tách riêng theo từng người để mình gửi trao đổi; đây là đề nghị phối hợp, không phải đánh giá hay xác nhận công việc của các bạn.')
heading('Kết quả đã có')
bullets([
    'Plugin kleer-plugin đã kích hoạt. Sau đồng bộ develop, giữ registry pa_* của nhóm và chờ WooCommerce để tạo thuộc tính; chưa thể kiểm chứng menu khi chưa có WooCommerce và sản phẩm.',
    '298 kiểm tra đạt, gồm các ca v4; toàn bộ 384 tổ hợp v1 chạy trên hai catalog kiểm thử. Giữ v4 6 câu của develop; PO cần chốt phiên bản dùng nghiệm thu.',
    'G3 chưa được duyệt. API hợp lệ hiện trả 503 vì thiếu WooCommerce; mẫu đo HTTP khoảng 0,40 đến 0,40 giây. Chưa đạt nghiệm thu API 200 hoặc dưới 300 ms.',
])
heading('Điều kiện còn thiếu theo task')
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [Cm(3.0), Cm(8.0), Cm(6.0)]
for column, width in zip(table.columns, widths): column.width = width
for cell, title in zip(table.rows[0].cells, ['Task của Nhân','Đầu vào hoặc phê duyệt còn thiếu','Phối hợp chính']): cell.text = title
rows = [
    ('L01 G3 01','Duyệt ma trận và đủ 8 tiêu chí DoR','Phát và An'),
    ('L01 G4 01','G3 được duyệt rồi review prompt v2','Phát và An'),
    ('L01 G5 03','PO chốt ma trận và cách đặt mã hồ sơ','Phát'),
    ('L08 G5 02','WooCommerce và kiểm tra menu thực tế','An và QA'),
    ('L01 G5 05','Dữ liệu thật và Postman 200 dưới 300 ms','Trí An Tuyền QA'),
    ('L01 G5 06','Fallback thật đủ bước và ngoại lệ 503','Trí Phát QA'),
    ('L09 G5 06','Baseline SQL EXPLAIN và số đo HTTP','An Trí QA'),
]
for row in rows:
    for index, (cell, text) in enumerate(zip(table.add_row().cells, row)):
        cell.text = text.replace(' ', '-') if index == 0 else text
for row_index, row in enumerate(table.rows):
    for col_index, cell in enumerate(row.cells):
        cell.width = widths[col_index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tc_pr = cell._tc.get_or_add_tcPr()
        margins = OxmlElement('w:tcMar')
        for side in ['top','left','bottom','right']:
            item = OxmlElement('w:'+side); item.set(qn('w:w'),'85'); item.set(qn('w:type'),'dxa'); margins.append(item)
        tc_pr.append(margins)
        if row_index == 0:
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'),'E8EEF4'); tc_pr.append(shade)
        elif row_index % 2 == 0:
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'),'F6F7F8'); tc_pr.append(shade)
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.line_spacing = 1.0
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
                run.bold = row_index == 0
    tr_pr = row._tr.get_or_add_trPr()
    no_split = OxmlElement('w:cantSplit'); tr_pr.append(no_split)
tbl_pr = table._tbl.tblPr
borders = OxmlElement('w:tblBorders')
for side in ['top','left','bottom','right','insideH','insideV']:
    b = OxmlElement('w:'+side); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'5'); b.set(qn('w:color'),'D9D9D9'); borders.append(b)
tbl_pr.append(borders)
repeat = OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
note = p('L09-G1-01 đã hoàn thành từ trước. Đợt này có thêm bằng chứng kích hoạt plugin L01-G5-01 và gói tài liệu L10-G8-01. Những task phụ thuộc phê duyệt hoặc dữ liệu thật chưa được coi là hoàn thành.')
note.paragraph_format.space_before = Pt(8)

people = [
    {
        'name':'Nguyễn Tấn Phát', 'role':'PO và BA',
        'intro':'Phát ơi, mình đã chuẩn bị bản kỹ thuật L01 theo Flow Card hiện có. Nhờ bạn duyệt các quyết định nghiệp vụ dưới đây để mình có đầu vào chính thức, vì G3 hiện chưa được duyệt.',
        'tasks':'L01-G3-01, L01-G4-01, L01-G5-03 và L01-G5-06',
        'needs':[
            'Chốt bản dùng nghiệm thu là v1 5 câu hay v4 6 câu; xác nhận từng lựa chọn và từng trọng số. Bản code đã bỏ hydration +1 ở flaky_rough để khớp ma trận hiện có; sensitivity >= 3 được coi là nhạy cảm.',
            'Chốt quy tắc bằng điểm. Đề xuất hiện tại ưu tiên oily, dry, combination, normal theo thứ tự khai báo. Nếu cần quy tắc khác, cho mình bản cụ thể để sửa và kiểm thử lại.',
            'Chốt mã hồ sơ nhạy cảm. Bản chuẩn bị dùng <loại_da>_sensitive, ví dụ dry_sensitive, để khớp ca nghiệm thu 2; vấn đề ưu tiên vẫn hiện trong nhãn và concerns.',
            'Duyệt sản phẩm fallback dịu nhẹ cho từng bước và quyết định ngoại lệ thiếu kho. Đề xuất: khi kho không thể đáp ứng an toàn, trả JSON 503 thay vì đưa ID mẫu hoặc routine thiếu.',
            'Duyệt sơ đồ hiện hành và các ca nghiệm thu; cùng An xác nhận đủ 8 tiêu chí DoR trước khi mình đóng G3.',
        ],
        'files':'Flow Card và approval.md tại docs/flows/L01-quiz-routine/; prompt v2 tại docs/prompts/PRM-L01-01-quiz-engine.md; đặc tả docs/api/openapi.json.',
        'evidence':'Bản ma trận và sơ đồ có số phiên bản/đường dẫn; từng quyết định chấp thuận hoặc yêu cầu sửa; xác nhận G3 kèm ngày. Mình không ghi thay chữ ký hoặc coi unit test là phê duyệt nghiệp vụ.',
        'reply':'Phát phản hồi giúp mình: bản ma trận được chốt; quy tắc bằng điểm; mã hồ sơ nhạy cảm; danh sách fallback; có chấp nhận 503 khi thiếu kho không; các mục còn phải sửa trước khi duyệt G3.',
    },
    {
        'name':'Lê Nguyễn Quốc An', 'role':'Scrum Master và DevOps',
        'intro':'An ơi, kleer-plugin đã được kích hoạt thành công ở local và phần tài liệu API đã chuẩn bị. Nhờ bạn review kỹ thuật và hỗ trợ điều kiện môi trường để mình nghiệm thu các task còn lại.',
        'tasks':'L01-G3-01, L08-G5-02, L01-G5-05 và L09-G5-06',
        'needs':[
            'Cùng Phát xác nhận G3 và người cập nhật 8 tiêu chí DoR sau khi duyệt. Hồ sơ hiện đã chuyển về chờ phê duyệt, không còn ghi đã ký thay các bạn.',
            'Cung cấp hoặc hướng dẫn môi trường nhóm có WooCommerce. Docker local hiện chỉ có WordPress, 0 sản phẩm và chưa có product post type; plugin của mình không tự tạo product type thay WooCommerce.',
            'Review thay đổi backend trên develop theo cách nhóm đã thống nhất push trực tiếp. Giữ nguyên v4 và giao diện mới của nhóm. Không đưa secret hoặc .env lên Git.',
            'Phối hợp đo toàn bộ HTTP trên môi trường đại diện. Hai request hiện trả 503 và mất khoảng 0,40 đến 0,40 giây; không thể coi đó là đạt mục tiêu dưới 300 ms.',
            'Xác nhận baseline 15 truy vấn trong DoD L09-G5-06 lấy từ đâu và phạm vi đo. Adapter mới có 4 SELECT theo lô trong kiểm thử; chưa có bằng chứng giảm 15 xuống 4 trên dữ liệu thật. Nếu cần điều chỉnh DoD thì nhờ nhóm quyết định.',
        ],
        'files':'docs/tasks/NHAN_NEXT_STEPS.md; docs/evidence/nhan/runtime.json; docs/evidence/nhan/http-benchmark.json; docs/api/README.md.',
        'evidence':'URL môi trường thử, phiên bản WooCommerce và dữ liệu; kết quả review; điều kiện benchmark cold/warm; baseline SQL hoặc quyết định điều chỉnh chỉ tiêu. Gửi quyền truy cập qua kênh an toàn nếu cần, không ghi mật khẩu vào tài liệu này.',
        'reply':'An cho mình biết: môi trường nào dùng nghiệm thu; khi nào có WooCommerce và kho; người review code; cách chốt G3; baseline 15 query và cách đo 300 ms.',
    },
    {
        'name':'Nguyễn Minh Trí', 'role':'Phụ trách L08 và Plugin Dev 2',
        'intro':'Trí ơi, backend quiz đã có bộ đọc sản phẩm thật và hỗ trợ cả taxonomy chuẩn lẫn thuộc tính pa_*. Nhờ bạn bàn giao đầu vào dữ liệu L08 để mình thử routine và fallback trên đúng ID WooCommerce.',
        'tasks':'L08-G5-02, L01-G5-05, L01-G5-06 và L09-G5-06 của Nhân',
        'needs':[
            'Bàn giao kho đã import và danh sách ID WooCommerce thật. Cần biết môi trường, số sản phẩm và phiên bản dữ liệu; không dùng số thứ tự CSV làm ID sản phẩm.',
            'Gán đúng slug theo thuộc tính pa_skin_type gồm oily, dry, combination, normal, sensitive, all_skin_types; pa_skin_concern gồm acne, dark_spots, hydration, aging, soothing; pa_routine_step gồm cleanser, treatment, moisturizer, sunscreen.',
            'Với v1, mỗi simple product có một routine_step, đã publish, còn hàng và có giá hợp lệ; v4 giữ thêm các bước tẩy trang, toner, serum theo bản develop. Bộ đọc cũng nhận pa_skin_type, pa_skin_concern, pa_routine_step nếu dữ liệu đang dùng thuộc tính WooCommerce; nhờ xác nhận bộ slug thực tế trước khi nhập.',
            'Cùng Phát xác nhận ít nhất một sản phẩm fallback dịu nhẹ thật cho mỗi bước. Gán _kleer_gentle_fallback=yes hoặc đồng thời all_skin_types và soothing sau khi có xác nhận, không tự suy diễn độ dịu nhẹ từ tên.',
            'Cùng An/QA cung cấp tập dữ liệu đại diện để đo SQL và API. Các ca thiếu kho cần một bản sao dữ liệu riêng; mình không sửa tồn kho chung để chạy test.',
        ],
        'files':'docs/api/README.md phần Đầu vào sản phẩm cần Trí bàn giao; src/Repositories/RoutineProductRepository.php trong kleer-plugin; Postman collection trong docs/api/.',
        'evidence':'File hoặc bảng ID thật kèm các term, giá/tồn kho; 4 ID fallback và người duyệt; môi trường thử và thời điểm có thể bàn giao. Đây là đầu vào cần phối hợp, không phải nhận xét về tiến độ phần việc của bạn.',
        'reply':'Trí gửi giúp mình: môi trường/kho hiện có; bộ slug đã dùng; danh sách ID kiểm thử; 4 ID fallback đã được PO duyệt; thời điểm dữ liệu sẵn sàng.',
    },
    {
        'name':'Huỳnh Ngọc Nhã Uyên', 'role':'Theme Dev 2 phụ trách Quiz Frontend và Cart Checkout UX',
        'intro':'Uyên ơi, mình đã chuẩn bị contract API và xử lý lỗi cho quiz. Nhờ bạn phối hợp sơ đồ và phần gọi API ở giao diện để mình đối chiếu backend L01; phần sửa theme vẫn do bạn phụ trách.',
        'tasks':'Đầu vào G3 và tích hợp cho L01-G5-03, L01-G5-05, L01-G5-06 của Nhân',
        'needs':[
            'Gửi đường dẫn phiên bản hiện hành của L01_Quiz_Logic_Flowchart.drawio và các màn hình quiz/routine liên quan để Phát duyệt cùng Flow Card.',
            'Đối chiếu question_id/option theo bản PO chốt: v1 có 5 câu, v4 có 6 câu và query version=v4. Giữ giao diện v4 mới nhất trên develop. Nếu UI hiện dùng tên khác, gửi bảng ánh xạ để mình cùng bạn chốt, không đoán từ nhãn hiển thị.',
            'Thống nhất POST /kleer/v1/skin-quiz hoặc alias /recommend, payload có session UUID và consent. Khi đang đăng nhập, gửi X-WP-Nonce hợp lệ; client bên thứ ba cần Origin đã được duyệt.',
            'Hiển thị routine theo thứ tự và ID thật API trả về; xử lý 400 bằng vị trí lỗi, 403 bằng làm mới nonce/kiểm tra Origin, 503 bằng thông báo và nút thử lại. Không vẽ kết quả mẫu như thể là routine thật khi API lỗi.',
            'Sau khi có dữ liệu của Trí, cùng mình thử ca 3 bước, 4 bước và thiếu kho; kiểm tra dữ liệu ID bàn giao cho chức năng thêm trọn bộ vào giỏ.',
        ],
        'files':'docs/api/openapi.json; docs/api/README.md; tools/quiz/valid-submission.json; docs/api/KLEER-Routine.postman_collection.json.',
        'evidence':'Link sơ đồ/frame có phiên bản; bảng ánh xạ câu hỏi; request/response từ giao diện trên môi trường thử; ảnh hoặc video ca thành công và lỗi. Các phần này cần thực hiện sau khi PO chốt nghiệp vụ.',
        'reply':'Uyên gửi giúp mình link sơ đồ, bảng tên câu hỏi/lựa chọn ở frontend và thời điểm có thể cùng thử API; cho mình biết chỗ nào trong contract cần thay để khớp giao diện.',
    },
    {
        'name':'Trần Ngọc Minh Tuyền', 'role':'Integration Dev phụ trách REST API và các tích hợp',
        'intro':'Tuyền ơi, mình đã có OpenAPI và Postman cho endpoint routine. Nhờ bạn review phần hợp đồng tích hợp và bảo mật để client của nhóm dùng thống nhất với backend của mình.',
        'tasks':'Review hợp đồng cho L01-G5-05 và bàn giao tài liệu L10-G8-01 của Nhân',
        'needs':[
            'Kiểm tra namespace kleer/v1, hai route /skin-quiz và /recommend, request/response schema, mã profile dự thảo và cách bàn giao ID sang luồng L03.',
            'Xác nhận host/cổng dùng nghiệm thu và danh sách Origin cần cho phép. Khi permalink mặc định, collection dùng /?rest_route=/kleer/v1/recommend; không bắt buộc sửa cấu hình website chỉ để gọi thử.',
            'Đối chiếu nonce cho browser có cookie đăng nhập và server client không có cookie. Client gửi nonce sai phải nhận 403; không chia sẻ khóa/secret trong tài liệu hoặc commit.',
            'Review các lỗi 400/403/503/500 và cách bên thứ ba xử lý thiếu kho. Import OpenAPI/Postman để kiểm tra tính sử dụng của tài liệu, rồi gửi các điểm cần sửa.',
            'Phối hợp An và Nhã Uyên kiểm tra request thật sau khi WooCommerce/dữ liệu đã có. API 503 hiện tại là thiếu dependency, không phải bằng chứng đã đạt tích hợp.',
        ],
        'files':'docs/api/openapi.json; docs/api/KLEER-Routine.postman_collection.json; docs/api/README.md; docs/evidence/nhan/http-contract-check.json.',
        'evidence':'Kết quả import và review; host/Origin được chấp thuận; ví dụ request đúng từ client; danh sách thay đổi hợp đồng cần phối hợp. Mọi thay đổi schema dùng chung cần được thống nhất trước.',
        'reply':'Tuyền phản hồi giúp mình: contract có điểm nào chưa khớp; môi trường và Origin nào dùng; nonce/client xử lý thế nào; tài liệu cần bổ sung gì cho bên tích hợp.',
    },
    {
        'name':'Nguyễn Ngọc Thảo Uyên', 'role':'QA và Automation',
        'intro':'Uyên ơi, phần backend của mình đã có kiểm thử độc lập và collection API. Nhờ bạn phối hợp lập bằng chứng nghiệm thu sau khi Phát/An duyệt G3 và Trí bàn giao dữ liệu; mình không thay kết quả unit test bằng kết luận QA.',
        'tasks':'Nghiệm thu L08-G5-02, L01-G5-03, L01-G5-05, L01-G5-06 và L09-G5-06 của Nhân',
        'needs':[
            'Đối chiếu test với ma trận PO thực sự duyệt, nhất là ca dry_sensitive, trọng số flaky_rough và xử lý bằng điểm. Bản hiện tại kiểm tra 384 tổ hợp với hai catalog fixture, chưa có dữ liệu thật.',
            'Trên môi trường có WooCommerce, kiểm tra 3 menu taxonomy, giá/tồn kho, ID sản phẩm trả về và routine đúng 3 hoặc 4 bước qua Postman/giao diện.',
            'Thử thiếu exact match nhưng có fallback, thiếu một bước, kho rỗng và sản phẩm không tương thích trên bản sao dữ liệu. Xác nhận 503 theo ngoại lệ PO chấp nhận và không có routine thiếu/ID mẫu.',
            'Chạy các ca payload thiếu/sai, câu hỏi trùng, consent false, Origin không được phép và nonce sai; lưu request, response và kết quả mong đợi. Không ghi session/cookie thật vào tài liệu gửi rộng.',
            'Cùng An đo cold/warm HTTP và số query trên kho đại diện; ghi điều kiện đo, mẫu đầu, trung vị và lớn nhất. HTTP lỗi 503 hoặc thời gian service đơn lẻ không chứng minh đạt dưới 300 ms.',
        ],
        'files':'wp-content/plugins/kleer-plugin/tests/run_nhan_tests.php; tools/quiz/verify-runtime.php; tools/quiz/benchmark.py; docs/api/ và docs/evidence/nhan/.',
        'evidence':'Test run có phiên bản code/dữ liệu, expected/actual, ảnh hoặc response JSON, số đo HTTP/SQL và kết luận theo DoD. Ghi rõ case chưa chạy do thiếu dependency thay vì đánh dấu PASS.',
        'reply':'Uyên cho mình biết bộ ca cần bổ sung, thời điểm có thể chạy cùng, môi trường/test data yêu cầu và mẫu báo cáo QA nhóm đang dùng.',
    },
    {
        'name':'Lâm Trúc Uyên', 'role':'Theme Dev 1 phụ trách Catalog và UI',
        'intro':'Uyên ơi, mình cần thống nhất dữ liệu taxonomy mà backend quiz đọc với bộ lọc Catalog. Đây là đề nghị kiểm tra tương thích đầu vào, không phải nhận xét hay yêu cầu mình sửa phần theme của bạn.',
        'tasks':'Phối hợp tương thích taxonomy cho L08-G5-02 và L01-G5-05 của Nhân',
        'needs':[
            'Cho mình biết Catalog đang lọc theo skin_type, skin_concern, routine_step hay các thuộc tính pa_* tương ứng. Backend đã đọc được cả hai nhóm tên nếu term dùng đúng slug.',
            'Đối chiếu với Trí bộ slug và một số ID sản phẩm mẫu. Nhãn tiếng Việt có thể dùng để hiển thị, nhưng request và dữ liệu lọc cần thống nhất slug để quiz không bỏ sót sản phẩm.',
            'Sau khi môi trường có WooCommerce và dữ liệu thật, cùng thử một sản phẩm xuất hiện đúng trong Catalog và trong bước routine tương ứng. Nếu cần đổi contract, mình cùng bạn và Tuyền/Trí chốt trước khi sửa.',
        ],
        'files':'docs/api/README.md phần bảng taxonomy và giới hạn sản phẩm; docs/api/openapi.json.',
        'evidence':'Tên taxonomy/filter đang dùng, bảng slug nếu có khác biệt và ví dụ ID thật; phản hồi các điểm không tương thích. Không cần bạn xác nhận những task ngoài phần phối hợp này.',
        'reply':'Uyên gửi giúp mình tên taxonomy/bộ lọc đang dùng, các slug khác với bảng backend và thời điểm tiện đối chiếu một lượt trên dữ liệu thật.',
    },
]

for person in people:
    doc.add_page_break()
    doc.add_heading(person['name'], level=1)
    p(person['role'])
    p(person['intro'])
    p('Liên quan phần việc của Nhân: ' + person['tasks'] + '.')
    heading('Việc cần phối hợp')
    bullets(person['needs'])
    heading('Tài liệu mình đã chuẩn bị')
    p(person['files'])
    heading('Đầu ra cần nhận để đối chiếu')
    p(person['evidence'])
    heading('Nội dung nhờ bạn phản hồi')
    p(person['reply'])
    p('Thời điểm có thể hỗ trợ hoặc bàn giao: ____________________')

footer = section.footer.paragraphs[0]
footer.add_run('KLEER  |  Nhân đề nghị phối hợp  |  ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
for run in footer.runs: run.font.size = Pt(9)
doc.save(output)
print(output.name)
